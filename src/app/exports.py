"""Artifact exports: the Reports flow metrics as .xlsx or .pdf, plus a generic
table→xlsx used by copilot artifacts.

Pure-Python only (openpyxl + fpdf2) — safe on the Databricks Apps runtime.
weasyprint was rejected deliberately: it needs system cairo/pango, and a
failing pip install would block every deploy.

Every duration in these artifacts is 5-day-workweek business days, matching
the tab and the copilot — one unit everywhere.
"""
from datetime import datetime, timezone
from io import BytesIO
from typing import Any


def _sheet(wb, title: str, headers: list[str], rows: list[list[Any]], widths: list[int]) -> None:
    from openpyxl.styles import Font

    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.freeze_panes = "A2"


def _date_sk(sk: Any) -> str:
    s = str(sk)
    return f"{s[:4]}-{s[4:6]}-{s[6:8]}" if len(s) == 8 else s


def reports_workbook(project: str, p: dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()

    # -- Summary -----------------------------------------------------------------
    ws = wb.active
    ws.title = "Summary"
    k = p["kpis"]
    f = p.get("filters", {})
    rows = [
        ("Project", project),
        ("Range", p["range"]),
        ("Generated (UTC)", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M")),
        ("Duration unit", "business days (Mon–Fri)"),
        ("Type filter", ", ".join(f.get("types", [])) or "all"),
        ("Assignee filter", ", ".join(f.get("assignees", [])) or "all"),
        (),
        ("Completed", k["throughput"]),
        ("Created", k["created"]),
        ("Net flow (created − completed)", k["netFlow"]),
        ("WIP now", k["wip"]),
        ("Cycle time p50 (bdays)", k["cycleP50"]),
        ("Cycle time p85 (bdays)", k["cycleP85"]),
        ("Oldest WIP (bdays)", k["oldestWipDays"]),
    ]
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_col=1, max_col=1):
        row[0].font = Font(bold=True)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 24

    # -- Created vs completed ------------------------------------------------------
    by_date: dict[str, list[int]] = {}
    for r in p["createdPerDay"]:
        by_date.setdefault(_date_sk(r["dateSK"]), [0, 0])[0] = r["count"]
    for r in p["completedPerDay"]:
        by_date.setdefault(_date_sk(r["dateSK"]), [0, 0])[1] = r["count"]
    _sheet(
        wb, "Created vs completed",
        ["Date", "Created", "Completed"],
        [[d, c[0], c[1]] for d, c in sorted(by_date.items())],
        [14, 10, 12],
    )

    # -- Cycle times ---------------------------------------------------------------
    _sheet(
        wb, "Cycle times",
        ["ID", "Title", "Type", "Closed", "Cycle (bdays)", "Lead (bdays)"],
        [[i["id"], i["title"], i["type"], i["closedDate"], i["cycleBdays"], i["leadBdays"]]
         for i in p["cycleItems"]],
        [8, 50, 12, 12, 14, 14],
    )

    # -- Open items (aging WIP) ------------------------------------------------------
    _sheet(
        wb, "Open items",
        ["ID", "Title", "Type", "State", "Category", "Assignee", "Age (bdays)"],
        [[i["id"], i["title"], i["type"], i["state"], i["category"],
          i.get("assignee") or "Unassigned", i["ageBdays"]]
         for i in sorted(p["openItems"], key=lambda x: -x["ageBdays"])],
        [8, 50, 12, 14, 12, 22, 12],
    )

    # -- Cumulative flow --------------------------------------------------------------
    _sheet(
        wb, "Cumulative flow",
        ["Date", "State category", "Open items"],
        [[r["date"], r["category"], r["count"]] for r in p["cfd"]],
        [14, 16, 12],
    )

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def table_workbook(name: str, columns: list[str], rows: list[list[Any]]) -> bytes:
    """One-sheet workbook for arbitrary tabular data (copilot/Genie artifacts)."""
    from openpyxl import Workbook

    # sheet titles reject \ / * ? : [ ] and empty — names come from user questions
    title = "".join(c for c in (name or "") if c not in "\\/*?:[]")[:31].strip() or "data"
    wb = Workbook()
    wb.remove(wb.active)
    widths = [max(12, min(50, len(str(c)) + 4)) for c in columns]
    _sheet(wb, title, columns, rows, widths)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---- PDF (fpdf2) --------------------------------------------------------------------

def _latin(s: Any) -> str:
    """Core PDF fonts are latin-1; degrade anything else instead of crashing."""
    return str(s if s is not None else "—").encode("latin-1", "replace").decode("latin-1")


def _pdf_table(pdf, headers: list[str], rows: list[list[Any]], widths: list[int]) -> None:
    from fpdf.enums import TableCellFillMode

    pdf.set_font("Helvetica", size=8)
    with pdf.table(
        col_widths=widths,
        text_align="LEFT",
        borders_layout="HORIZONTAL_LINES",
        cell_fill_color=(243, 244, 246),
        cell_fill_mode=TableCellFillMode.ROWS,
        line_height=5.5,
    ) as table:
        head = table.row()
        pdf.set_font("Helvetica", style="B", size=8)
        for h in headers:
            head.cell(_latin(h))
        pdf.set_font("Helvetica", size=8)
        for r in rows:
            row = table.row()
            for v in r:
                row.cell(_latin(v))


def reports_pdf(project: str, p: dict[str, Any]) -> bytes:
    from fpdf import FPDF

    k = p["kpis"]
    f = p.get("filters", {})
    pdf = FPDF(orientation="P", format="A4")
    pdf.set_auto_page_break(auto=True, margin=14)
    pdf.add_page()

    pdf.set_font("Helvetica", style="B", size=16)
    pdf.cell(0, 9, _latin(f"Flow metrics - {project}"), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", size=9)
    pdf.set_text_color(110)
    sub = (
        f"Range {p['range']}  |  generated {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M')} UTC"
        f"  |  durations in business days (Mon-Fri)"
        f"  |  types: {', '.join(f.get('types', [])) or 'all'}"
        f"  |  assignees: {', '.join(f.get('assignees', [])) or 'all'}"
    )
    pdf.cell(0, 6, _latin(sub), new_x="LMARGIN", new_y="NEXT")
    pdf.set_text_color(0)
    pdf.ln(3)

    # KPI strip
    kpis = [
        ("Completed", k["throughput"]), ("Created", k["created"]),
        ("Net flow", k["netFlow"]), ("WIP now", k["wip"]),
        ("Cycle p50", f"{k['cycleP50']}d"), ("Cycle p85", f"{k['cycleP85']}d"),
        ("Oldest WIP", f"{k['oldestWipDays']}d"),
    ]
    cell_w = (pdf.w - pdf.l_margin - pdf.r_margin) / len(kpis)
    pdf.set_font("Helvetica", size=7)
    pdf.set_text_color(110)
    for label, _ in kpis:
        pdf.cell(cell_w, 4, _latin(label.upper()), align="C")
    pdf.ln()
    pdf.set_font("Helvetica", style="B", size=13)
    pdf.set_text_color(0)
    for _, value in kpis:
        pdf.cell(cell_w, 8, _latin(value), align="C")
    pdf.ln(12)

    def section(title: str) -> None:
        pdf.set_font("Helvetica", style="B", size=11)
        pdf.cell(0, 8, _latin(title), new_x="LMARGIN", new_y="NEXT")
        pdf.ln(1)

    by_date: dict[str, list[int]] = {}
    for r in p["createdPerDay"]:
        by_date.setdefault(_date_sk(r["dateSK"]), [0, 0])[0] = r["count"]
    for r in p["completedPerDay"]:
        by_date.setdefault(_date_sk(r["dateSK"]), [0, 0])[1] = r["count"]
    if by_date:
        section("Created vs completed (per day)")
        _pdf_table(pdf, ["Date", "Created", "Completed"],
                   [[d, c[0], c[1]] for d, c in sorted(by_date.items())], [30, 20, 20])
        pdf.ln(6)

    if p["cycleItems"]:
        section("Cycle times (completed items)")
        _pdf_table(
            pdf, ["ID", "Title", "Type", "Closed", "Cycle (bd)", "Lead (bd)"],
            [[i["id"], i["title"][:60], i["type"], i["closedDate"], i["cycleBdays"], i["leadBdays"]]
             for i in p["cycleItems"]],
            [12, 70, 18, 22, 18, 18],
        )
        pdf.ln(6)

    if p["openItems"]:
        section("Open items (aging WIP, oldest first)")
        _pdf_table(
            pdf, ["ID", "Title", "State", "Assignee", "Age (bd)"],
            [[i["id"], i["title"][:60], i["state"], i.get("assignee") or "Unassigned", i["ageBdays"]]
             for i in sorted(p["openItems"], key=lambda x: -x["ageBdays"])],
            [12, 80, 22, 34, 16],
        )

    return bytes(pdf.output())
