"""Reports .xlsx export: sheet structure and values round-trip via openpyxl."""
from io import BytesIO

from openpyxl import load_workbook

from app.exports import reports_pdf, reports_workbook, table_workbook

PAYLOAD = {
    "range": "30d",
    "days": 30,
    "durationUnit": "businessDays",
    "filters": {"types": ["Issue"], "assignees": []},
    "kpis": {
        "throughput": 2, "created": 3, "netFlow": 1, "wip": 2,
        "cycleP50": 2.0, "cycleP85": 5.0, "oldestWipDays": 9,
    },
    "createdPerDay": [{"dateSK": 20260701, "count": 2}, {"dateSK": 20260702, "count": 1}],
    "completedPerDay": [{"dateSK": 20260702, "count": 2}],
    "cycleItems": [
        {"id": 4, "title": "Fix", "type": "Issue", "closedDate": "2026-07-02",
         "cycleBdays": 2, "leadBdays": 3},
    ],
    "openItems": [
        {"id": 5, "title": "Old one", "type": "Issue", "state": "Doing",
         "category": "InProgress", "assignee": None, "ageBdays": 9},
        {"id": 6, "title": "New one", "type": "Issue", "state": "To Do",
         "category": "Proposed", "assignee": "Ada", "ageBdays": 1},
    ],
    "cfd": [{"date": "2026-07-01", "category": "Proposed", "count": 3}],
}


def test_workbook_sheets_and_values():
    wb = load_workbook(BytesIO(reports_workbook("home", PAYLOAD)))
    assert wb.sheetnames == [
        "Summary", "Created vs completed", "Cycle times", "Open items", "Cumulative flow",
    ]

    s = wb["Summary"]
    kv = {r[0].value: r[1].value for r in s.iter_rows() if r[0].value}
    assert kv["Project"] == "home"
    assert kv["Type filter"] == "Issue"
    assert kv["Cycle time p85 (bdays)"] == 5.0
    assert "business days" in kv["Duration unit"]

    cvc = wb["Created vs completed"]
    rows = [[c.value for c in r] for r in cvc.iter_rows(min_row=2)]
    assert rows == [["2026-07-01", 2, 0], ["2026-07-02", 1, 2]]  # merged by date

    open_ws = wb["Open items"]
    first = [c.value for c in open_ws[2]]
    assert first[0] == 5 and first[5] == "Unassigned" and first[6] == 9  # oldest first

    assert [c.value for c in wb["Cycle times"][2]] == [4, "Fix", "Issue", "2026-07-02", 2, 3]
    assert [c.value for c in wb["Cumulative flow"][2]] == ["2026-07-01", "Proposed", 3]


def test_reports_pdf_renders():
    data = reports_pdf("home", PAYLOAD)
    assert data[:5] == b"%PDF-"
    assert len(data) > 1200  # KPIs + three tables, not an empty shell


def test_table_workbook_roundtrip():
    wb = load_workbook(BytesIO(table_workbook("items by state", ["state", "n"], [["To Do", 3]])))
    ws = wb["items by state"]
    assert [c.value for c in ws[1]] == ["state", "n"]
    assert [c.value for c in ws[2]] == ["To Do", 3]
